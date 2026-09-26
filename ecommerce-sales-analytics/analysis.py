"""E-Commerce Sales & Customer Analytics.

Pipeline: load raw CSVs -> clean -> build fact tables -> KPIs, trends, category,
channel, RFM segmentation, cohort retention, Pareto, discounts -> SQL checks
-> charts (images/), Excel workbook (outputs/), HTML dashboard (docs/).
Run:  python analysis.py
"""
import json
import sqlite3
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).parent
RAW, IMG, OUT, DOCS = ROOT / "data/raw", ROOT / "images", ROOT / "outputs", ROOT / "docs"
for p in (IMG, OUT, OUT / "sql", DOCS):
    p.mkdir(exist_ok=True)

# =============================================================================
# 1. Load & clean
# =============================================================================
customers = pd.read_csv(RAW / "customers.csv", parse_dates=["signup_date"])
products = pd.read_csv(RAW / "products.csv")
orders = pd.read_csv(RAW / "orders.csv")
items = pd.read_csv(RAW / "order_items.csv")

quality = []


def note(issue, table, rows, fix):
    quality.append({"Issue": issue, "Table": table, "Rows affected": int(rows), "Fix": fix})


n = len(orders)
orders = orders.drop_duplicates()
note("Exact duplicate rows", "orders", n - len(orders), "Dropped")

odd = ~orders["order_date"].str.fullmatch(r"\d{4}-\d{2}-\d{2}")
orders["order_date"] = pd.to_datetime(orders["order_date"], format="mixed")
note("Non-ISO date format (e.g. 05-Mar-2024)", "orders", odd.sum(), "Parsed to ISO dates")

raw_region = customers["region"]
customers["region"] = raw_region.str.strip().str.title()
note("Inconsistent region casing/whitespace", "customers",
     (raw_region.notna() & (raw_region != customers["region"])).sum(), "Trimmed + title-cased")
note("Missing region", "customers", customers["region"].isna().sum(), "Labelled 'Unknown'")
customers["region"] = customers["region"].fillna("Unknown")

bad_qty = items["quantity"] <= 0
note("Zero/negative quantity", "order_items", bad_qty.sum(), "Dropped (invalid lines)")
items = items[~bad_qty]

missing_price = items["unit_price"].isna()
items.loc[missing_price, "unit_price"] = items.loc[missing_price, "product_id"].map(
    products.set_index("product_id")["unit_price"])
note("Missing unit price", "order_items", missing_price.sum(), "Filled from product catalog")

cancelled = orders["status"].eq("Cancelled").sum()
note("Cancelled orders", "orders", cancelled, "Excluded from sales metrics")
quality = pd.DataFrame(quality)
print(quality.to_string(index=False), "\n")

# =============================================================================
# 2. Fact tables
# =============================================================================
df = (items.merge(orders[orders["status"] == "Completed"], on="order_id")
      .merge(products[["product_id", "product_name", "category", "unit_cost"]], on="product_id")
      .merge(customers, on="customer_id"))
df["gross_sales"] = df["quantity"] * df["unit_price"] * (1 - df["discount_pct"] / 100)
df["net_sales"] = df["gross_sales"].where(~df["returned"], 0)          # returns refunded
df["cogs"] = (df["quantity"] * df["unit_cost"]).where(~df["returned"], 0)
df["gross_profit"] = df["net_sales"] - df["cogs"]

ords = (df.groupby(["order_id", "customer_id", "order_date", "discount_pct"], as_index=False)
        [["gross_sales", "net_sales", "gross_profit", "quantity"]].sum()
        .sort_values("order_date"))
ords["year"] = ords["order_date"].dt.year
ords["month"] = ords["order_date"].dt.to_period("M")
ords["customer_type"] = np.where(ords.groupby("customer_id").cumcount() == 0, "New", "Returning")

cust = (ords.groupby("customer_id")
        .agg(orders=("order_id", "size"), net_sales=("net_sales", "sum"),
             first_order=("order_date", "min"), last_order=("order_date", "max"))
        .join(customers.set_index("customer_id")))

# =============================================================================
# 3. KPIs (overall + YoY)
# =============================================================================


def kpis(o):
    return {
        "Net Sales": o["net_sales"].sum(),
        "Gross Profit": o["gross_profit"].sum(),
        "Gross Margin %": o["gross_profit"].sum() / o["net_sales"].sum(),
        "Orders": len(o),
        "Customers": o["customer_id"].nunique(),
        "Avg Order Value": o["net_sales"].mean(),
        "Orders per Customer": len(o) / o["customer_id"].nunique(),
        "Repeat Customer Rate": (o.groupby("customer_id").size() > 1).mean(),
        "Return Rate (% of sales)": 1 - o["net_sales"].sum() / o["gross_sales"].sum(),
    }


kpi = pd.DataFrame({"All": kpis(ords), "2023": kpis(ords[ords.year == 2023]), "2024": kpis(ords[ords.year == 2024])})
kpi["YoY Change"] = kpi["2024"] / kpi["2023"] - 1
print(kpi.round(3), "\n")

# =============================================================================
# 4. Monthly trend
# =============================================================================
monthly = ords.groupby("month").agg(net_sales=("net_sales", "sum"), orders=("order_id", "size"),
                                    customers=("customer_id", "nunique"), gross_profit=("gross_profit", "sum"))
monthly = monthly.join(ords.pivot_table(index="month", columns="customer_type", values="net_sales", aggfunc="sum")
                       .rename(columns=lambda c: f"{c.lower()}_customer_sales"))
monthly["aov"] = monthly["net_sales"] / monthly["orders"]
monthly["yoy_growth"] = monthly["net_sales"].pct_change(12)
monthly.index = monthly.index.astype(str)

# =============================================================================
# 5. Category, channel, region
# =============================================================================
category = df.groupby("category").agg(net_sales=("net_sales", "sum"), gross_sales=("gross_sales", "sum"),
                                      gross_profit=("gross_profit", "sum"), units=("quantity", "sum"),
                                      orders=("order_id", "nunique"))
category["sales_share"] = category["net_sales"] / category["net_sales"].sum()
category["gross_margin"] = category["gross_profit"] / category["net_sales"]
category["return_rate"] = 1 - category["net_sales"] / category["gross_sales"]
category["profit_share"] = category["gross_profit"] / category["gross_profit"].sum()
category = category.sort_values("net_sales", ascending=False)

top_products = (df.groupby(["product_id", "product_name", "category"], as_index=False)
                .agg(net_sales=("net_sales", "sum"), units=("quantity", "sum"), gross_profit=("gross_profit", "sum"))
                .nlargest(15, "net_sales"))

# 12-month customer value: only customers whose first order is >= 365 days before data end (fair comparison)
data_end = ords["order_date"].max()
o = ords.merge(cust[["first_order"]], left_on="customer_id", right_index=True)
ltv12 = o[o["order_date"] < o["first_order"] + pd.Timedelta(days=365)].groupby("customer_id")["net_sales"].sum()
cust["ltv_12m"] = ltv12.where(cust["first_order"] <= data_end - pd.Timedelta(days=365))

channel = cust.groupby("acquisition_channel").agg(
    customers=("orders", "size"), net_sales=("net_sales", "sum"),
    avg_orders=("orders", "mean"), repeat_rate=("orders", lambda s: (s > 1).mean()),
    ltv_12m=("ltv_12m", "mean"))
channel["sales_share"] = channel["net_sales"] / channel["net_sales"].sum()
channel = channel.sort_values("ltv_12m", ascending=False)

region = cust.groupby("region").agg(customers=("orders", "size"), net_sales=("net_sales", "sum"),
                                    sales_per_customer=("net_sales", "mean"),
                                    repeat_rate=("orders", lambda s: (s > 1).mean()))
region["sales_share"] = region["net_sales"] / region["net_sales"].sum()
region = region.sort_values("net_sales", ascending=False)

# =============================================================================
# 6. RFM segmentation
# =============================================================================
snapshot = data_end + pd.Timedelta(days=1)
rfm = cust[["acquisition_channel", "region"]].copy()
rfm["recency_days"] = (snapshot - cust["last_order"]).dt.days
rfm["frequency"] = cust["orders"]
rfm["monetary"] = cust["net_sales"]
rfm["R"] = pd.qcut(rfm["recency_days"], 5, labels=[5, 4, 3, 2, 1]).astype(int)
# Most customers order once, so quantiles on frequency are meaningless -> fixed, explainable buckets
rfm["F"] = pd.cut(rfm["frequency"], [0, 1, 2, 3, 5, np.inf], labels=[1, 2, 3, 4, 5]).astype(int)
rfm["M"] = pd.qcut(rfm["monetary"].rank(method="first"), 5, labels=[1, 2, 3, 4, 5]).astype(int)
fm = (rfm["F"] + rfm["M"]) / 2
R = rfm["R"]
rfm["segment"] = np.select(
    [(R >= 4) & (fm >= 4), (R >= 3) & (fm >= 3), R >= 4, R == 3, (R <= 2) & (fm >= 3), R == 2],
    ["Champions", "Loyal", "New / Promising", "Needs Attention", "At Risk", "Hibernating"],
    default="Lost")
seg_order = ["Champions", "Loyal", "New / Promising", "Needs Attention", "At Risk", "Hibernating", "Lost"]
segments = rfm.groupby("segment").agg(customers=("R", "size"), net_sales=("monetary", "sum"),
                                      avg_recency_days=("recency_days", "mean"),
                                      avg_orders=("frequency", "mean"), avg_sales=("monetary", "mean")).reindex(seg_order)
segments["customer_share"] = segments["customers"] / segments["customers"].sum()
segments["sales_share"] = segments["net_sales"] / segments["net_sales"].sum()

# =============================================================================
# 7. Cohort retention (by month of first order)
# =============================================================================
mi = ords["order_date"].dt.year * 12 + ords["order_date"].dt.month
ords["cohort"] = ords.groupby("customer_id")["month"].transform("min").astype(str)
ords["months_since_first"] = mi - mi.groupby(ords["customer_id"]).transform("min")
cohort_counts = ords.pivot_table(index="cohort", columns="months_since_first", values="customer_id", aggfunc="nunique")
# no buyers = 0, but months after the data ends stay blank
cohort_mi = pd.PeriodIndex(cohort_counts.index, freq="M")
observable = (data_end.year * 12 + data_end.month) - (cohort_mi.year * 12 + cohort_mi.month)
cohort_counts = cohort_counts.fillna(0).where(cohort_counts.columns.to_numpy() <= observable.to_numpy()[:, None])
retention = cohort_counts.div(cohort_counts[0], axis=0)
retention.insert(0, "cohort_size", cohort_counts[0])
avg_retention = retention.drop(columns="cohort_size").mean()

# =============================================================================
# 8. Pareto & discounts
# =============================================================================
pareto = cust["net_sales"].sort_values(ascending=False).cumsum() / cust["net_sales"].sum()
top20_share = pareto.iloc[int(len(pareto) * 0.2) - 1]
top10_share = pareto.iloc[int(len(pareto) * 0.1) - 1]

discount = ords.groupby("discount_pct").agg(orders=("order_id", "size"), aov=("net_sales", "mean"),
                                            units_per_order=("quantity", "mean"),
                                            net_sales=("net_sales", "sum"), gross_profit=("gross_profit", "sum"))
discount["gross_margin"] = discount["gross_profit"] / discount["net_sales"]
discount["profit_per_order"] = discount["gross_profit"] / discount["orders"]
discount.index = [f"{d}%" for d in discount.index]
discount.index.name = "discount"

print(f"Top 10% customers = {top10_share:.1%} of sales, top 20% = {top20_share:.1%}")
print(category.round(3), "\n", channel.round(3), "\n", segments.round(2), "\n", discount.round(3), "\n")
print("Avg retention by month:", avg_retention.head(13).round(3).to_dict(), "\n")

# =============================================================================
# 9. SQL: same questions answered in SQLite (sql/analysis_queries.sql)
# =============================================================================
con = sqlite3.connect(":memory:")
customers.to_sql("customers", con, index=False)
products.to_sql("products", con, index=False)
orders.assign(order_date=orders["order_date"].dt.strftime("%Y-%m-%d")).to_sql("orders", con, index=False)
items.assign(returned=items["returned"].astype(int)).to_sql("order_items", con, index=False)
for q in (s for s in (ROOT / "sql/analysis_queries.sql").read_text().split(";") if s.strip()):
    title = q.strip().splitlines()[0].lstrip("- ").strip()
    if "CREATE VIEW" in q:
        con.execute(q)
        continue
    res = pd.read_sql(q, con)
    res.to_csv(OUT / "sql" / f"{title.split()[0].lower()}.csv", index=False)
    print(f"SQL {title}\n{res.head(5).to_string(index=False)}\n")
# cross-check: SQL and pandas must agree on total net sales
sql_total = pd.read_sql("SELECT SUM(net_sales) s FROM (SELECT * FROM v_line_sales)", con)["s"][0]
assert abs(sql_total - ords["net_sales"].sum()) < 1, (sql_total, ords["net_sales"].sum())

# =============================================================================
# 10. Charts
# =============================================================================
BLUE, ORANGE, GREY, INK = "#2563eb", "#f59e0b", "#94a3b8", "#1e293b"
plt.rcParams.update({"figure.dpi": 130, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.titleweight": "bold", "axes.titlesize": 13, "axes.titlelocation": "left",
                     "font.size": 10, "axes.edgecolor": GREY, "text.color": INK, "axes.labelcolor": INK})
k = lambda v, _=None: f"${v / 1000:,.0f}K"  # noqa: E731


def save(fig, name):
    fig.tight_layout()
    fig.savefig(IMG / name, bbox_inches="tight")
    plt.close(fig)


fig, ax = plt.subplots(figsize=(11, 4.5))
x = np.arange(len(monthly))
ax.bar(x, monthly["new_customer_sales"], color=BLUE, label="New customers")
ax.bar(x, monthly["returning_customer_sales"], bottom=monthly["new_customer_sales"], color=ORANGE, label="Returning customers")
ax.set_xticks(x, monthly.index, rotation=60, ha="right")
ax.yaxis.set_major_formatter(k)
ax.set_title("Monthly net sales: holiday peaks and a growing share from returning customers")
ax.legend(frameon=False)
save(fig, "01_monthly_sales.png")

fig, ax = plt.subplots(figsize=(9, 4.5))
c = category.sort_values("net_sales")
ax.barh(c.index, c["net_sales"], color=BLUE)
for y, (s, m, r) in enumerate(zip(c["net_sales"], c["gross_margin"], c["return_rate"])):
    ax.text(s, y, f"  margin {m:.0%} · returns {r:.0%}", va="center", fontsize=9)
ax.xaxis.set_major_formatter(k)
ax.set_xlim(0, c["net_sales"].max() * 1.45)
ax.set_title("Net sales by category (with gross margin and return rate)")
save(fig, "02_category_performance.png")

fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4.2))
a1.bar(channel.index, channel["ltv_12m"], color=BLUE)
a1.set_title("Avg 12-month customer value")
a1.yaxis.set_major_formatter(lambda v, _: f"${v:,.0f}")
a2.bar(channel.index, channel["repeat_rate"], color=ORANGE)
a2.set_title("Repeat-purchase rate")
a2.yaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
for a in (a1, a2):
    a.tick_params(axis="x", rotation=25)
save(fig, "03_channel_value.png")

fig, ax = plt.subplots(figsize=(10, 4.5))
x = np.arange(len(segments))
ax.bar(x - 0.2, segments["customer_share"], 0.4, color=GREY, label="% of customers")
ax.bar(x + 0.2, segments["sales_share"], 0.4, color=BLUE, label="% of net sales")
ax.set_xticks(x, segments.index, rotation=15)
ax.yaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
ax.set_title("RFM segments: share of customers vs share of sales")
ax.legend(frameon=False)
save(fig, "04_rfm_segments.png")

heat = retention.drop(columns="cohort_size").iloc[:12, :13]
fig, ax = plt.subplots(figsize=(11, 5.5))
ax.imshow(heat.iloc[:, 1:].values, cmap="Blues", vmin=0, vmax=0.25, aspect="auto")
for (i, j), v in np.ndenumerate(heat.iloc[:, 1:].values):
    if not np.isnan(v):
        ax.text(j, i, f"{v:.0%}", ha="center", va="center", fontsize=8, color="white" if v > 0.15 else INK)
ax.set_xticks(range(12), [f"M{m}" for m in range(1, 13)])
ax.set_yticks(range(len(heat)), [f"{p} (n={n:.0f})" for p, n in zip(heat.index, retention["cohort_size"].iloc[:12])])
ax.set_title("Cohort retention: % of each 2023 cohort ordering again N months later")
for s in ax.spines.values():
    s.set_visible(False)
save(fig, "05_cohort_retention.png")

fig, ax = plt.subplots(figsize=(8, 4.5))
xs = np.arange(1, len(pareto) + 1) / len(pareto)
ax.plot(xs, pareto.values, color=BLUE, lw=2.5)
ax.plot([0, 1], [0, 1], color=GREY, ls="--", lw=1)
ax.axvline(0.2, color=ORANGE, lw=1)
ax.annotate(f"Top 20% of customers\n= {top20_share:.0%} of net sales", (0.2, top20_share),
            xytext=(0.35, top20_share - 0.2), arrowprops={"arrowstyle": "->", "color": INK})
ax.xaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
ax.yaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
ax.set_xlabel("Customers (ranked by net sales)")
ax.set_ylabel("Cumulative share of net sales")
ax.set_title("Customer concentration (Pareto curve)")
save(fig, "06_pareto.png")

fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4))
a1.bar(discount.index, discount["units_per_order"], color=BLUE)
a1.set_title("Units per order by discount level")
a2.bar(discount.index, discount["gross_margin"], color=ORANGE)
a2.yaxis.set_major_formatter(lambda v, _: f"{v:.0%}")
a2.set_title("Gross margin by discount level")
save(fig, "07_discount_impact.png")

# =============================================================================
# 11. Excel workbook with native-chart dashboard
# =============================================================================
from openpyxl.chart import BarChart, LineChart, Reference  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

PCT = ("share", "margin", "rate", "growth", "yoy", "%")
MONEY = ("sales", "profit", "aov", "ltv", "monetary", "value")
xlsx = OUT / "ecommerce_analytics.xlsx"
sheets = {
    "KPIs": kpi, "Monthly Trend": monthly, "Categories": category, "Top Products": top_products,
    "Channels": channel, "Regions": region, "RFM Segments": segments, "Cohort Retention": retention,
    "Discounts": discount, "Customer RFM": rfm.sort_values("monetary", ascending=False), "Data Quality": quality,
}
with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
    pd.DataFrame().to_excel(xw, sheet_name="Dashboard")
    for name, t in sheets.items():
        t.to_excel(xw, sheet_name=name, index=name not in ("Top Products", "Data Quality"))
    wb = xw.book
    head_fill, head_font = PatternFill("solid", fgColor="1E293B"), Font(color="FFFFFF", bold=True)

    for name in sheets:
        ws = wb[name]
        ws.freeze_panes = "B2"
        for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
            header = str(col[0].value or "").lower()
            col[0].fill, col[0].font = head_fill, head_font
            fmt = ("0.0%" if any(p in header for p in PCT) or name == "Cohort Retention" and header.isdigit()
                   else "$#,##0" if any(m in header for m in MONEY) else None)
            for cell in col[1:]:
                if fmt and isinstance(cell.value, (int, float)):
                    cell.number_format = fmt
            width = max(len(str(c.value)) for c in col[:200] if c.value is not None) if any(c.value for c in col) else 8
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 45)
    # KPIs sheet: row-wise formats
    ws = wb["KPIs"]
    for row in ws.iter_rows(min_row=2):
        label = row[0].value
        fmt = "0.0%" if "%" in label or "Rate" in label else "0.00" if "per" in label else "#,##0" if label in ("Orders", "Customers") else "$#,##0"
        for c in row[1:4]:
            c.number_format = fmt
        row[4].number_format = "+0.0%;-0.0%"
    ws = wb["Cohort Retention"]
    ws.conditional_formatting.add(f"D2:{get_column_letter(ws.max_column)}{ws.max_row}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                 end_type="num", end_value=0.25, end_color="2563EB"))

    # ---- Dashboard sheet
    db = wb["Dashboard"]
    db.sheet_view.showGridLines = False
    db["B2"] = "E-Commerce Sales & Customer Analytics Dashboard (Jan 2023 – Dec 2024)"
    db["B2"].font = Font(size=18, bold=True, color="1E293B")
    tiles = [("Net Sales", "Net Sales", "$#,##0"), ("Orders", "Orders", "#,##0"),
             ("Customers", "Customers", "#,##0"), ("Avg Order Value", "Avg Order Value", "$#,##0.00"),
             ("Gross Margin", "Gross Margin %", "0.0%"), ("Repeat Rate", "Repeat Customer Rate", "0.0%")]
    kpi_rows = {label: i + 2 for i, label in enumerate(kpi.index)}
    for i, (title, key, fmt) in enumerate(tiles):
        col = get_column_letter(2 + i * 3)
        db[f"{col}4"] = title
        db[f"{col}5"] = f"=KPIs!B{kpi_rows[key]}"
        db[f"{col}6"] = f'="2024 vs 2023: "&TEXT(KPIs!E{kpi_rows[key]},"+0.0%;-0.0%")'
        db[f"{col}4"].font = Font(size=10, color="64748B", bold=True)
        db[f"{col}5"].font = Font(size=20, bold=True, color="2563EB")
        db[f"{col}5"].number_format = fmt
        db[f"{col}6"].font = Font(size=9, color="64748B")
        for r in (4, 5, 6):
            db[f"{col}{r}"].fill = PatternFill("solid", fgColor="F1F5F9")
            db[f"{col}{r}"].alignment = Alignment(horizontal="left")
    for c in range(2, 20):
        db.column_dimensions[get_column_letter(c)].width = 11

    def chart(kind, title, sheet, data_cols, anchor, rows, cats_col=1, w=17, h=8, **kw):
        ch = kind()
        ch.title, ch.width, ch.height, ch.style = title, w, h, 10
        ws = wb[sheet]
        for dc in data_cols:
            ch.add_data(Reference(ws, min_col=dc, min_row=1, max_row=rows + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=cats_col, min_row=2, max_row=rows + 1))
        for k_, v in kw.items():
            setattr(ch, k_, v)
        db.add_chart(ch, anchor)
        return ch

    mcols = list(monthly.columns)
    ch = chart(BarChart, "Monthly Net Sales: New vs Returning Customers", "Monthly Trend",
               [mcols.index("new_customer_sales") + 2, mcols.index("returning_customer_sales") + 2],
               "B8", len(monthly), w=34, grouping="stacked", overlap=100)
    ch.y_axis.numFmt = "$#,##0"
    chart(BarChart, "Net Sales by Category", "Categories", [2], "B25", len(category), type="bar")
    chart(BarChart, "Avg 12-Month Customer Value by Channel", "Channels", [list(channel.columns).index("ltv_12m") + 2],
          "K25", len(channel))
    ch = chart(BarChart, "RFM Segments: % Customers vs % Sales", "RFM Segments",
               [list(segments.columns).index("customer_share") + 2, list(segments.columns).index("sales_share") + 2],
               "B42", len(segments))
    ch.y_axis.numFmt = "0%"
    # retention line chart from a helper row of average retention under the cohort table
    ws = wb["Cohort Retention"]
    base = ws.max_row + 3
    ws.cell(base, 3, "Average retention")  # columns D.. hold months 1..12, same as the table above
    for j, v in enumerate(avg_retention.iloc[1:13], start=1):
        ws.cell(base - 1, j + 3, f"M{j}")
        ws.cell(base, j + 3, float(v)).number_format = "0.0%"
    ch = LineChart()
    ch.title, ch.width, ch.height, ch.style = "Avg Retention by Months Since First Order", 17, 8, 10
    ch.add_data(Reference(ws, min_col=3, max_col=15, min_row=base), from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=4, max_col=15, min_row=base - 1))
    ch.y_axis.numFmt = "0%"
    db.add_chart(ch, "K42")
    wb.move_sheet("Dashboard", offset=-wb.index(wb["Dashboard"]))
    wb.active = 0

# =============================================================================
# 12. HTML dashboard data (docs/ is served by GitHub Pages)
# =============================================================================
r = lambda s: s.round(4).tolist()  # noqa: E731
payload = {
    "kpi": {k_: {"all": float(v["All"]), "yoy": float(v["YoY Change"])} for k_, v in kpi.iterrows()},
    "monthly": {"labels": list(monthly.index), "new": r(monthly["new_customer_sales"]),
                "returning": r(monthly["returning_customer_sales"]), "aov": r(monthly["aov"])},
    "category": {"labels": list(category.index), "sales": r(category["net_sales"]),
                 "margin": r(category["gross_margin"]), "returns": r(category["return_rate"])},
    "channel": {"labels": list(channel.index), "ltv": r(channel["ltv_12m"]), "repeat": r(channel["repeat_rate"]),
                "customers": channel["customers"].tolist()},
    "segments": {"labels": seg_order, "customers": r(segments["customer_share"]), "sales": r(segments["sales_share"]),
                 "count": segments["customers"].tolist()},
    "retention": {"cohorts": list(retention.index), "size": retention["cohort_size"].astype(int).tolist(),
                  "rows": [[None if pd.isna(v) else round(float(v), 4) for v in row]
                           for row in retention.drop(columns="cohort_size").iloc[:, 1:13].values]},
    "discount": {"labels": list(discount.index), "units": r(discount["units_per_order"]),
                 "margin": r(discount["gross_margin"]), "profit": r(discount["profit_per_order"])},
    "pareto": {"top10": round(float(top10_share), 4), "top20": round(float(top20_share), 4)},
}
(DOCS / "data.js").write_text("window.DASHBOARD_DATA = " + json.dumps(payload) + ";\n")
print(f"Done -> {xlsx.relative_to(ROOT)}, images/, docs/data.js")
